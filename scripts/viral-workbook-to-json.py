"""Read-only Excel conversion. Requires openpyxl; no workbook is edited."""
import argparse,datetime,hashlib,json,pathlib
import openpyxl
p=argparse.ArgumentParser(description='Prepare a private migration JSON. No Supabase/Drive connection.')
p.add_argument('xlsx');p.add_argument('output_json');p.add_argument('--bill-folder-id',default='');p.add_argument('--private-xlsx');a=p.parse_args()
def extract(path):
 values=openpyxl.load_workbook(path,data_only=True);formula_book=openpyxl.load_workbook(path,data_only=False);sheets={};formulas=[]
 for s in values:
  headers=[str(c.value or '') for c in s[1]];rows=[headers]
  for row in s.iter_rows(min_row=2):
   if not any(c.value is not None and c.value!='' for c in row):continue
   output=[]
   for i,c in enumerate(row):
    v=c.value;k=headers[i]
    if isinstance(v,datetime.datetime):v=v.date().isoformat() if v.hour==v.minute==v.second==0 else v.isoformat()+'+07:00'
    elif isinstance(v,datetime.date):v=v.isoformat()
    elif isinstance(v,datetime.time):v=v.strftime('%H:%M')
    elif isinstance(v,(int,float)) and not isinstance(v,bool) and (k.endswith('_id') or k=='bank_account' or k=='account_number'):
     if abs(v)>=2**53:raise ValueError('Numeric ID/account exceeds safe precision at '+s.title+'!'+c.coordinate)
     v=str(int(v)) if float(v).is_integer() else str(v)
    output.append('' if v is None else v)
   rows.append(output)
  for row in formula_book[s.title]:
   for c in row:
    if c.data_type=='f':formulas.append({'sheet':s.title,'cell':c.coordinate,'formula':c.value if isinstance(c.value,str) else getattr(c.value,'text',str(c.value)),'cached':str(values[s.title][c.coordinate].value or '')})
  sheets[s.title]=rows
 return sheets,formulas
src=pathlib.Path(a.xlsx);sheets,formulas=extract(src)
result={'format':'viral-workbook-v1','fileName':src.name,'sourceSHA256':hashlib.sha256(src.read_bytes()).hexdigest(),'exportedAt':datetime.datetime.now(datetime.timezone.utc).isoformat(),'sheets':sheets,'formulas':formulas,'driveFolderId':a.bill_folder_id}
if a.private_xlsx:
 private=pathlib.Path(a.private_xlsx);result['privateSheets'],result['privateFormulas']=extract(private);result['privateSourceSHA256']=hashlib.sha256(private.read_bytes()).hexdigest()
out=pathlib.Path(a.output_json);out.parent.mkdir(parents=True,exist_ok=True);out.write_text(json.dumps(result,ensure_ascii=False,allow_nan=False),encoding='utf8')
print('Prepared',len(sheets),'sheets;',out.stat().st_size,'bytes. Keep this JSON private, not in GitHub.')
